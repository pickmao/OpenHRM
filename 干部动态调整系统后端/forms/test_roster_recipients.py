from datetime import timedelta
from io import BytesIO
import tempfile
from unittest import skipUnless
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from accounts.models import Role, User, UserRole
from cadres.models import PersonnelRoster
from orgs.models import Membership, OrgUnit
from .models import FormTask, OnlyOfficeDocument


class DispatchRosterRecipientTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create(username='roster-admin', is_superuser=True)
        self.client.force_authenticate(self.admin)

    def person(self, name, department='一部门', **kwargs):
        return PersonnelRoster.objects.create(
            serial_number=PersonnelRoster.all_objects.count() + 1,
            name=name, department=department, gender='M', **kwargs,
        )

    def options(self):
        response = self.client.get('/api/forms/dispatch/options/')
        self.assertEqual(response.status_code, 200, response.data)
        return response.data['roster_people']

    def test_roster_people_without_accounts_are_visible_and_read_does_not_create_accounts(self):
        person = self.person('实际填报人')
        account_count = User.objects.count()
        rows = self.options()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['roster_id'], str(person.id))
        self.assertEqual(rows[0]['name'], person.name)
        self.assertEqual(rows[0]['department'], person.department)
        self.assertEqual(rows[0]['account_state'], 'MISSING')
        self.assertIsNone(rows[0]['user_id'])
        self.assertEqual(User.objects.count(), account_count)

    def test_unique_account_can_receive_roster_task_and_history_is_excluded(self):
        person = self.person('在册人员')
        self.person('历史人员', is_active=False)
        account = User.objects.create_user(username='actual', real_name=person.name)
        row, = self.options()
        self.assertEqual(row['account_state'], 'READY')
        self.assertEqual(row['user_id'], str(account.id))
        self.assertEqual(row['username'], 'actual')

    def test_disabled_account_is_not_treated_as_missing_or_usable(self):
        person = self.person('停用人员')
        User.objects.create_user(username='disabled', real_name=person.name, is_active=False)
        row, = self.options()
        self.assertEqual(row['account_state'], 'DISABLED')
        self.assertIsNone(row['user_id'])

    def test_same_name_accounts_require_unique_department_evidence(self):
        person = self.person('同名人员')
        one = User.objects.create_user(username='one', real_name=person.name)
        User.objects.create_user(username='two', real_name=person.name)
        row, = self.options()
        self.assertEqual(row['account_state'], 'AMBIGUOUS')
        unit = OrgUnit.objects.create(name=person.department, unit_type='DEPARTMENT')
        Membership.objects.create(user=one, unit=unit, is_primary=True)
        row, = self.options()
        self.assertEqual(row['user_id'], str(one.id))

    def test_one_account_is_not_assigned_to_two_same_name_roster_people(self):
        self.person('同名人员', '一部门')
        self.person('同名人员', '二部门')
        User.objects.create_user(username='one', real_name='同名人员')
        self.assertTrue(all(row['account_state'] == 'AMBIGUOUS' for row in self.options()))

    def test_roster_options_require_dispatch_permission(self):
        self.client.force_authenticate(User.objects.create_user(username='ordinary'))
        self.assertEqual(self.client.get('/api/forms/dispatch/options/').status_code, 403)
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get('/api/forms/dispatch/options/').status_code, 401)

    @skipUnless(connection.vendor == 'postgresql', '现有表单权限查询依赖 PostgreSQL JSON contains')
    def test_roster_recipient_can_login_fill_excel_submit_and_update_progress(self):
        person = self.person('实际填报人')
        account = User.objects.create_user(username='actual-filler', password='Roster-QA-2026!', real_name=person.name)
        role = Role.objects.create(code='FORM_FILLER', name='填报人', permissions=['forms:task:view', 'forms:task:submit'])
        UserRole.objects.create(user=account, role=role)
        other = User.objects.create_user(username='unassigned', password='Roster-QA-2026!')
        UserRole.objects.create(user=other, role=role)
        user_id = self.options()[0]['user_id']
        workbook = Workbook()
        workbook.active['A1'] = '请填写'
        original = BytesIO()
        workbook.save(original)
        with tempfile.TemporaryDirectory() as media, override_settings(MEDIA_ROOT=media):
            template = self.client.post('/api/forms/templates/', {
                'name': '实际人员填报表', 'code': 'ROSTER_FLOW',
                'source_file': SimpleUploadedFile('roster.xlsx', original.getvalue()),
            }, format='multipart')
            self.assertEqual(template.status_code, 201, template.data)
            payload = {'batch_name': '花名册下发测试', 'deadline_at': (timezone.now() + timedelta(days=1)).isoformat(),
                       'rules': [{'template_id': template.data['id'], 'receiver_type': 'USER',
                                  'receiver_expr_json': {'user_ids': [user_id]}}]}
            preview = self.client.post('/api/forms/dispatch/preview/', payload, format='json')
            self.assertEqual(preview.status_code, 200, preview.data)
            self.assertEqual(preview.data['summary']['total'], 1)
            published = self.client.post('/api/forms/dispatch/publish/', payload, format='json')
            self.assertEqual(published.status_code, 201, published.data)
            self.assertEqual(published.data['summary']['created_tasks'], 1)
            batch_id = published.data['batch']['id']
            self.client.force_authenticate(None)
            login = self.client.post('/api/auth/login/', {'username': account.username, 'password': 'Roster-QA-2026!'}, format='json')
            self.assertEqual(login.status_code, 200, login.data)
            self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {login.data["access"]}')
            tasks = self.client.get('/api/forms/tasks/my/')
            self.assertEqual(tasks.status_code, 200, tasks.data)
            self.assertEqual(len(tasks.data), 1)
            task_id = tasks.data[0]['id']
            config = self.client.get(f'/api/forms/tasks/{task_id}/onlyoffice-config/')
            self.assertEqual(config.status_code, 200, config.data)
            workbook.active['A1'] = '实际人员已填写'
            edited = BytesIO()
            workbook.save(edited)
            with patch('forms.views.urlopen', return_value=BytesIO(edited.getvalue())):
                callback = self.client.post(f'/api/forms/tasks/{task_id}/onlyoffice-callback/', {
                    'key': config.data['document']['key'], 'status': 6, 'url': 'https://document-server.example/filled.xlsx',
                }, format='json')
            self.assertEqual(callback.status_code, 200, callback.data)
            self.assertEqual(callback.data['error'], 0)
            submitted = self.client.post(f'/api/forms/tasks/{task_id}/submit/', {}, format='json')
            self.assertEqual(submitted.status_code, 200, submitted.data)
            self.assertEqual(FormTask.objects.get(id=task_id).status, 'SUBMITTED')
            document = OnlyOfficeDocument.objects.get(task_id=task_id)
            with document.file.open('rb') as saved:
                self.assertEqual(load_workbook(saved).active['A1'].value, '实际人员已填写')
            self.client.credentials()
            self.client.force_authenticate(other)
            self.assertEqual(self.client.get('/api/forms/tasks/my/').data, [])
            self.assertEqual(self.client.get(f'/api/forms/tasks/{task_id}/onlyoffice-config/').status_code, 403)
            self.client.force_authenticate(self.admin)
            progress = self.client.get(f'/api/forms/dispatch/{batch_id}/progress/')
            self.assertEqual(progress.status_code, 200, progress.data)
            self.assertEqual(progress.data['summary']['submitted'], 1)
