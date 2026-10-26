import tempfile
from datetime import timedelta
from io import BytesIO
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from accounts.models import User
from .models import DispatchBatch, FormTask, FormTemplate
from .services import create_onlyoffice_document


class TemplateUploadTests(TestCase):
    def setUp(self):
        self.media = tempfile.TemporaryDirectory()
        self.addCleanup(self.media.cleanup)
        settings = override_settings(MEDIA_ROOT=self.media.name)
        settings.enable()
        self.addCleanup(settings.disable)
        self.admin = User.objects.create(username='template-admin', is_superuser=True)
        self.client = APIClient()
        self.client.force_authenticate(self.admin)

    def excel_file(self):
        book = Workbook()
        book.active.append(['测试内容'])
        stream = BytesIO()
        book.save(stream)
        return SimpleUploadedFile('test.xlsx', stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def test_multipart_upload_is_enabled_and_can_be_previewed_and_published(self):
        response = self.client.post('/api/forms/templates/', {
            'name': '在线表格测试', 'code': 'UPLOAD_ENABLED', 'source_file': self.excel_file(),
        }, format='multipart')
        self.assertEqual(response.status_code, 201, response.data)
        self.assertTrue(response.data['is_active'])
        template = FormTemplate.objects.get(id=response.data['id'])
        self.assertTrue(template.is_active)
        self.assertTrue(template.source_file.storage.exists(template.source_file.name))
        payload = {'batch_name': '测试批次', 'deadline_at': (timezone.now() + timedelta(days=1)).isoformat(),
            'rules': [{'template_id': str(template.id), 'receiver_type': 'USER',
                'receiver_expr_json': {'user_ids': [str(self.admin.id)]}}]}
        response = self.client.post('/api/forms/dispatch/preview/', payload, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['summary']['total'], 1)
        response = self.client.post('/api/forms/dispatch/publish/', payload, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data['summary']['created_tasks'], 1)

    def test_explicit_disable_and_partial_metadata_update_keep_template_disabled(self):
        response = self.client.post('/api/forms/templates/', {
            'name': '停用模板', 'code': 'UPLOAD_DISABLED', 'is_active': 'false',
            'source_file': self.excel_file(),
        }, format='multipart')
        self.assertEqual(response.status_code, 201, response.data)
        self.assertFalse(response.data['is_active'])
        response = self.client.patch(f'/api/forms/templates/{response.data["id"]}/',
            {'name': '修改名称'}, format='multipart')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertFalse(response.data['is_active'])

    def test_repeated_force_saves_and_final_close_keep_all_updates(self):
        template = FormTemplate.objects.create(code='CALLBACK', name='连续保存', source_file=self.excel_file())
        batch = DispatchBatch.objects.create(name='保存测试', status='PUBLISHED',
            deadline_at=timezone.now() + timedelta(days=1))
        task = FormTask.objects.create(batch=batch, template=template,
            assignee_id=self.admin.id, deadline_at=batch.deadline_at)
        document = create_onlyoffice_document(task)
        key = document.document_key
        callback = f'/api/forms/tasks/{task.id}/onlyoffice-callback/'
        for index, (status, value) in enumerate([(6, '第一次保存'), (6, '第二次保存'), (2, '关闭前最终保存')], start=2):
            workbook = Workbook()
            workbook.active['A1'] = value
            saved = BytesIO()
            workbook.save(saved)
            with patch('forms.views.urlopen', return_value=BytesIO(saved.getvalue())):
                response = self.client.post(callback, {'key': key, 'status': status,
                    'url': 'https://document-server.example/saved.xlsx'}, format='json')
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual(response.data['error'], 0)
            document.refresh_from_db()
            self.assertEqual(document.version, index)
            self.assertEqual(document.last_callback_status, status)
            self.assertIsNotNone(document.last_saved_at)
            with document.file.open('rb') as stored:
                self.assertEqual(load_workbook(stored).active['A1'].value, value)
            if status == 6:
                self.assertEqual(document.document_key, key)
            else:
                self.assertNotEqual(document.document_key, key)
