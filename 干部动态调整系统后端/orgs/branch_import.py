"""支部-部门归属 Excel 导入。覆盖时只改 parent，不删除组织节点。"""
from __future__ import annotations

from io import BytesIO

import pandas as pd
from django.db import transaction
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .models import OrgUnit, UnitType

TEMPLATE_HEADERS = ['支部名称', '部门名称', '支部编码', '部门编码', '备注']
REQUIRED_HEADERS = ['支部名称', '部门名称']


def _normalize_header(value) -> str:
    text = '' if value is None or (isinstance(value, float) and pd.isna(value)) else str(value)
    return text.replace('\n', '').replace('\r', '').replace(' ', '').replace('*', '').strip()


def _cell_text(value) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ''
    text = str(value).strip()
    return '' if text.lower() == 'nan' else text


def build_column_map(columns) -> dict[str, str]:
    mapping = {}
    for col in columns:
        key = _normalize_header(col)
        if not key or key.startswith('Unnamed'):
            continue
        mapping.setdefault(key, str(col))
    return mapping


def missing_required_columns(column_map: dict[str, str]) -> list[str]:
    return [name for name in REQUIRED_HEADERS if name not in column_map]


def _row_text(row: pd.Series, column_map: dict[str, str], *names: str) -> str:
    for name in names:
        raw = column_map.get(name)
        if raw is None:
            continue
        text = _cell_text(row.get(raw))
        if text:
            return text
    return ''


def _index_units(units: list[OrgUnit]):
    by_code = {}
    by_name = {}
    for unit in units:
        if unit.code:
            by_code[unit.code] = unit
        by_name.setdefault(unit.name, []).append(unit)
    return by_code, by_name


def _lookup_unit(kind, name, code, by_code, by_name):
    if code:
        unit = by_code.get(code)
        if not unit:
            return None, f'找不到{kind}编码「{code}」'
        if name and unit.name != name:
            return None, f'{kind}编码「{code}」对应「{unit.name}」，与模板名称「{name}」不一致'
        return unit, None
    if not name:
        return None, f'{kind}名称和编码都为空'
    matches = by_name.get(name) or []
    if len(matches) == 1:
        return matches[0], None
    if len(matches) > 1:
        return None, f'{kind}名称「{name}」存在多个匹配，请填写编码'
    return None, f'找不到{kind}「{name}」'


def parse_branch_rows(df: pd.DataFrame) -> tuple[list[dict], list[dict]]:
    column_map = build_column_map(df.columns)
    missing = missing_required_columns(column_map)
    if missing:
        raise ValueError(f'Excel文件缺少必需的列: {", ".join(missing)}')

    rows = []
    errors = []
    for index, row in df.iterrows():
        excel_row = index + 2
        branch_name = _row_text(row, column_map, '支部名称')
        department_name = _row_text(row, column_map, '部门名称')
        branch_code = _row_text(row, column_map, '支部编码', '支部代码')
        department_code = _row_text(row, column_map, '部门编码', '部门代码')
        remark = _row_text(row, column_map, '备注')
        if not any([branch_name, department_name, branch_code, department_code]):
            continue
        if not (branch_name or branch_code) or not (department_name or department_code):
            errors.append({
                'row': excel_row,
                'branch_name': branch_name,
                'department_name': department_name,
                'error': '支部名称（或编码）和部门名称（或编码）均为必填',
            })
            continue
        rows.append({
            'row': excel_row,
            'branch_name': branch_name,
            'department_name': department_name,
            'branch_code': branch_code,
            'department_code': department_code,
            'remark': remark,
        })
    return rows, errors


def resolve_assignments(parsed_rows: list[dict]) -> tuple[list[dict], list[dict]]:
    branches = list(OrgUnit.objects.filter(unit_type=UnitType.BRANCH))
    departments = list(OrgUnit.objects.exclude(unit_type=UnitType.BRANCH))
    branch_by_code, branch_by_name = _index_units(branches)
    dept_by_code, dept_by_name = _index_units(departments)

    resolved = []
    errors = []
    seen_departments = {}

    for item in parsed_rows:
        branch, branch_error = _lookup_unit(
            '支部', item['branch_name'], item['branch_code'], branch_by_code, branch_by_name
        )
        department, dept_error = _lookup_unit(
            '部门', item['department_name'], item['department_code'], dept_by_code, dept_by_name
        )
        row_errors = [msg for msg in (branch_error, dept_error) if msg]
        if row_errors:
            errors.append({
                'row': item['row'],
                'branch_name': item['branch_name'],
                'department_name': item['department_name'],
                'error': '；'.join(row_errors),
            })
            continue

        previous = seen_departments.get(department.id)
        if previous:
            if previous != branch.id:
                errors.append({
                    'row': item['row'],
                    'branch_name': item['branch_name'],
                    'department_name': department.name,
                    'error': f'部门「{department.name}」在文件中归属多个支部',
                })
            continue
        seen_departments[department.id] = branch.id
        resolved.append({'branch': branch, 'department': department, 'remark': item['remark']})

    return resolved, errors


def apply_branch_assignments(resolved: list[dict], *, replace: bool) -> dict:
    assigned_ids = {item['department'].id for item in resolved}
    unassigned_count = 0
    assigned_count = 0

    with transaction.atomic():
        if replace:
            current_members = list(
                OrgUnit.objects.exclude(unit_type=UnitType.BRANCH).filter(
                    parent__unit_type=UnitType.BRANCH
                )
            )
            detach_ids = [unit.id for unit in current_members if unit.id not in assigned_ids]
            if detach_ids:
                unassigned_count = OrgUnit.objects.filter(id__in=detach_ids).update(parent=None)

        for item in resolved:
            department = item['department']
            branch = item['branch']
            if department.parent_id == branch.id:
                continue
            if branch in department.get_descendants():
                raise ValueError(f'不能将部门「{department.name}」纳入其下级支部「{branch.name}」')
            department.parent = branch
            department.save(update_fields=['parent', 'updated_at'])
            assigned_count += 1

    return {
        'assigned_count': assigned_count,
        'unassigned_count': unassigned_count,
        'mapped_count': len(resolved),
    }


def build_template_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '支部部门关系'
    sheet.append(TEMPLATE_HEADERS)
    for index in range(1, len(TEMPLATE_HEADERS) + 1):
        cell = sheet.cell(1, index)
        cell.font = Font(bold=True)
        sheet.column_dimensions[get_column_letter(index)].width = 18
    sheet.freeze_panes = 'A2'

    guide = workbook.create_sheet('填写说明')
    guide['A1'] = '覆盖导入说明'
    guide['A1'].font = Font(bold=True)
    lines = [
        '1. 必填列：支部名称、部门名称。可选列：支部编码、部门编码、备注。',
        '2. 上传并覆盖会按模板重建“部门归属哪个支部”：模板中的部门 parent 更新为对应支部。',
        '3. 当前已挂在支部下、但不在模板中的部门会移出支部（parent 置空），不会删除任何组织节点。',
        '4. 找不到的支部或部门会整表报错，不会部分写入。',
        '5. 一个部门同一时间只属于一个支部；文件内同一部门不能指向多个支部。',
        '6. 请先在系统中创建支部和部门，再通过本模板维护归属。',
    ]
    for offset, line in enumerate(lines, start=3):
        guide[f'A{offset}'] = line
    guide.column_dimensions['A'].width = 90

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
