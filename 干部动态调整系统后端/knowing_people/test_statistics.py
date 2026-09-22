from datetime import timedelta
from io import BytesIO

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from .models import InspectionCampaign, InspectionTask, TaskStatus
from .statistics import build_statistics, statistics_workbook


class StatisticsTests(TestCase):
    def setUp(self):
        self.campaign = InspectionCampaign.objects.create(name='统计验证', deadline_at=timezone.now() + timedelta(days=1))
        self.counter = 0

    def vote(self, role, grade='优', form='ATTACHMENT_7_1', group='chief', identity='1', name='同名',
             state=TaskStatus.SUBMITTED, recognition='认可', missing=None):
        self.counter += 1
        user = User.objects.create(username=f'stats{self.counter}')
        target = {'id': identity, 'name': name, 'target_group': group}
        scores = {key: grade for key in ['political', 'integrity', 'responsibility', 'style', 'thinking', 'frontline', 'performance']}
        if missing:
            scores.pop(missing)
        payload = {'targets': [{**target, 'scores': scores, 'recognition': recognition}]}
        if form == 'ATTACHMENT_6_2':
            payload = {'scores': scores, 'recognition': recognition}
        return InspectionTask.objects.create(campaign=self.campaign, assignee=user, form_type=form,
            filler_role=role, status=state, assignee_name_snapshot=user.username,
            context_json={'targets': [target]}, payload_json=payload)

    def rows(self, key):
        return next(s['rows'] for s in build_statistics(self.campaign)['report_sections'] if s['key'] == key)

    def test_branch_weighted_score_and_actual_denominator(self):
        self.vote('PRISON_LEADER', '优', 'ATTACHMENT_6_1', 'branch')
        self.vote('BRANCH_SECRETARY', '良', 'ATTACHMENT_6_1', 'branch')
        self.vote('BRANCH_STAFF', '中', 'ATTACHMENT_6_2', 'branch')
        self.vote('BRANCH_STAFF', '差', 'ATTACHMENT_6_2', 'branch')
        row = self.rows('branch')[-1]
        self.assertEqual(row['final_score'], 82.0)  # 95*.4 + 85*.2 + 67.5*.4
        self.assertEqual(row['effective_count'], 28)

    def test_chief_weights_and_committee_role_merging(self):
        for role, grade in [('PRINCIPAL_LEADER', '优'), ('PRISON_LEADER', '良'), ('BRANCH_SECRETARY', '中')]:
            self.vote(role, grade)
        self.vote('BRANCH_SECRETARY', '良', 'ATTACHMENT_7_4')
        self.vote('BRANCH_DEPUTY_SECRETARY', '优', 'ATTACHMENT_7_4')
        self.vote('BRANCH_STAFF', '差', 'ATTACHMENT_7_4')
        self.assertEqual(self.rows('chief')[-1]['final_score'], 80.2)

    def test_team_and_police_explicit_source_weights(self):
        for role, grade in [('POLITICAL_DIRECTOR', '优'), ('POLITICAL_EXECUTIVE_DEPUTY', '良'), ('POLITICAL_DEPUTY', '中')]:
            self.vote(role, grade, 'ATTACHMENT_7_3', 'team')
        for group in ['team', 'police']:
            for role, grade in [('BRANCH_SECRETARY', '优'), ('BRANCH_DEPUTY_SECRETARY', '良'), ('BRANCH_COMMITTEE', '中'), ('BRANCH_STAFF', '差')]:
                self.vote(role, grade, 'ATTACHMENT_7_4', group)
        self.assertEqual(self.rows('team')[-1]['final_score'], 79.6)
        self.assertEqual(self.rows('police')[-1]['final_score'], 73.8)

    def test_same_name_ids_and_non_submitted_exclusion(self):
        self.vote('BRANCH_SECRETARY', identity='a')
        self.vote('BRANCH_SECRETARY', identity='b')
        self.vote('BRANCH_SECRETARY', identity='c', state=TaskStatus.DRAFT)
        self.vote('BRANCH_SECRETARY', identity='d', state=TaskStatus.RETURNED)
        report = build_statistics(self.campaign)
        self.assertEqual(report['submitted_count'], 2)
        self.assertEqual({r['target_id'] for r in self.rows('chief') if r['effective_count']}, {'a', 'b'})
        self.assertEqual({r['target_id'] for r in self.rows('chief') if r['pending_count']}, {'c', 'd'})
        self.assertTrue(all(r['final_score'] is None for r in self.rows('chief')))

    def test_missing_dimension_blocks_overall_and_recognition_unknown_is_75(self):
        for role in ['PRINCIPAL_LEADER', 'PRISON_LEADER', 'BRANCH_SECRETARY']:
            self.vote(role, missing='performance' if role == 'PRISON_LEADER' else None)
        self.vote('BRANCH_COMMITTEE', form='ATTACHMENT_7_4', recognition='不了解')
        self.vote('BRANCH_STAFF', form='ATTACHMENT_7_4', recognition='不了解')
        self.assertIsNone(self.rows('chief')[-1]['final_score'])
        self.assertEqual(self.rows('chief')[-1]['missing_count'], 1)
        self.assertEqual(self.rows('recognition')[0]['final_score'], 75)

    def test_export_counts_and_formula_injection(self):
        self.vote('BRANCH_STAFF', form='ATTACHMENT_7_4', name='=1+1')
        workbook = load_workbook(BytesIO(statistics_workbook(self.campaign)))
        self.assertEqual(len(workbook.sheetnames), 9)
        self.assertEqual(workbook['正职统计'].max_row, 9)
        self.assertEqual(workbook['正职统计']['B2'].value, "'=1+1")
        self.assertEqual(workbook['认可度']['H2'].value, 95)

    def test_legacy_missing_ids_are_not_merged_by_name(self):
        self.vote('BRANCH_STAFF', identity=None)
        self.vote('BRANCH_STAFF', identity=None)
        identities = {r['target_id'] for r in self.rows('chief')}
        self.assertEqual(len(identities), 2)
        self.assertTrue(all(identity.startswith('legacy:') for identity in identities))

    def test_complete_groups_with_pending_or_returned_vote_have_no_final(self):
        for role in ['PRINCIPAL_LEADER', 'PRISON_LEADER', 'BRANCH_SECRETARY']:
            self.vote(role)
        self.vote('BRANCH_COMMITTEE', form='ATTACHMENT_7_4')
        self.vote('BRANCH_STAFF', form='ATTACHMENT_7_4')
        self.assertEqual(self.rows('chief')[-1]['final_score'], 95)
        pending = self.vote('BRANCH_STAFF', form='ATTACHMENT_7_4', state=TaskStatus.PENDING)
        # A draft must not hide its assigned target or add someone else's pending count.
        pending.payload_json = {'targets': [{'id': 'forged', 'name': '不应采用'}]}
        for state in [TaskStatus.PENDING, TaskStatus.DRAFT, TaskStatus.RETURNED]:
            with self.subTest(state=state):
                pending.status = state
                pending.save(update_fields=['status', 'payload_json'])
                rows = self.rows('chief')
                self.assertTrue(all(row['final_score'] is None for row in rows))
                self.assertTrue(all(row['pending_count'] == 1 for row in rows))
                self.assertEqual(rows[-1]['group_scores']['BRANCH_STAFF'], 95)
                self.assertIn('尚有1份未提交', rows[-1]['reason'])
                self.assertEqual({row['target_id'] for row in rows}, {'1'})
                self.assertIsNone(self.rows('recognition')[0]['final_score'])

    def test_merged_committee_label_and_fixed_branch_staff_source(self):
        self.vote('BRANCH_SECRETARY', form='ATTACHMENT_7_4')
        self.vote('BRANCH_SECRETARY', form='ATTACHMENT_6_2', group='branch')
        votes = build_statistics(self.campaign)['votes']
        committee = [row for row in votes if row['form_type'] == 'ATTACHMENT_7_4']
        staff = [row for row in votes if row['form_type'] == 'ATTACHMENT_6_2']
        self.assertTrue(all(row['filler_role'] == 'BRANCH_COMMITTEE' and row['filler_role_label'] == '党支部支委' for row in committee))
        self.assertTrue(all(row['filler_role'] == 'BRANCH_STAFF' for row in staff))
        self.assertNotIn('None', self.rows('branch')[0]['group_scores_text'])

    def test_missing_vote_does_not_hide_behind_other_valid_votes(self):
        for role in ['PRINCIPAL_LEADER', 'PRISON_LEADER', 'BRANCH_SECRETARY']:
            self.vote(role)
        self.vote('BRANCH_COMMITTEE', form='ATTACHMENT_7_4')
        self.vote('BRANCH_STAFF', form='ATTACHMENT_7_4')
        self.vote('BRANCH_STAFF', form='ATTACHMENT_7_4', missing='performance', recognition='')
        rows = self.rows('chief')
        self.assertEqual(rows[0]['final_score'], 95)
        self.assertIsNone(rows[-1]['final_score'])
        self.assertIsNone(next(row for row in rows if row['dimension_key'] == 'performance')['final_score'])
        self.assertIsNone(self.rows('recognition')[0]['final_score'])
