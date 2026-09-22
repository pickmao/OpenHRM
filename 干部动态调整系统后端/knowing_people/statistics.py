"""按原统计表计票；分母取实际有效票，缺少评分来源不生成最终分。"""
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_DOWN
from io import BytesIO

from .form_defs import TEAM_DIMENSIONS, form_catalog
from .models import FormType, FillerRole, TaskStatus

GRADES = {'优': 95, '良': 85, '中': 75, '差': 60}
RECOGNITION = {'认可': 95, '基本认可': 85, '不了解': 75, '不认可': 60}
SECTIONS = [('branch', '附件6 班子统计'), ('chief', '附件7 正职统计'),
            ('deputy', '附件7 副职统计'), ('team', '附件7 团队负责人统计'),
            ('police', '附件7 警员职级统计'), ('unclassified', '待确认对象类别'),
            ('recognition', '认可度统计')]
MATRIX_FORMS = {FormType.ATTACHMENT_6_1, FormType.ATTACHMENT_6_2,
                FormType.ATTACHMENT_7_1, FormType.ATTACHMENT_7_2,
                FormType.ATTACHMENT_7_3, FormType.ATTACHMENT_7_4}
NOTE = ('统计覆盖整个批次，不随任务列表筛选变化。只计已提交的有效评价；'
        '各维度按实际有效票人数平均，优/良/中/差为95/85/75/60分；'
        '最终分截取两位小数，仍有未提交任务或缺少来源、维度时不计最终分，不按零分或重新分配权重。')


def rounded(value):
    return float(value.quantize(Decimal('0.01'), rounding=ROUND_DOWN)) if value is not None else None


def _mean(values):
    return sum(map(Decimal, values)) / len(values) if values else None


def _targets(task):
    payload, context = task.payload_json or {}, task.context_json or {}
    snapshots = context.get('targets') or []
    if task.form_type == FormType.ATTACHMENT_6_2:
        snapshot = snapshots[0] if snapshots else context.get('branch') or {}
        yield {**snapshot, **payload, 'id': snapshot.get('id') or task.branch_id,
               'name': snapshot.get('name') or task.branch_name_snapshot}, snapshot, 0
        return
    lookup = {str(t.get('id')): t for t in snapshots if t.get('id') is not None}
    for index, target in enumerate(payload.get('targets') or []):
        yield target, lookup.get(str(target.get('id')), {}), index


def _section(task, snapshot):
    return {FormType.ATTACHMENT_6_1: 'branch', FormType.ATTACHMENT_6_2: 'branch',
            FormType.ATTACHMENT_7_1: 'chief', FormType.ATTACHMENT_7_2: 'deputy',
            FormType.ATTACHMENT_7_3: 'team'}.get(task.form_type, snapshot.get('target_group') or 'unclassified')


def _final(section, means):
    if section == 'branch':
        leaders = means.get('leaders')
        parts = [(leaders, Decimal('.4')), (means.get(FillerRole.BRANCH_SECRETARY), Decimal('.2')),
                 (means.get(FillerRole.BRANCH_STAFF), Decimal('.4'))]
    elif section in {'chief', 'deputy'}:
        parts = [(means.get(FillerRole.PRINCIPAL_LEADER), Decimal('.12')),
                 (means.get(FillerRole.PRISON_LEADER), Decimal('.28')),
                 (means.get(FillerRole.BRANCH_SECRETARY), Decimal('.2')),
                 (means.get(FillerRole.BRANCH_COMMITTEE), Decimal('.2')),
                 (means.get(FillerRole.BRANCH_STAFF), Decimal('.2'))]
    elif section in {'team', 'police'}:
        branch_weight = Decimal('.4') if section == 'team' else Decimal('.6')
        parts = [(means.get(FillerRole.BRANCH_SECRETARY), branch_weight * Decimal('.3')),
                 (means.get('BRANCH_DEPUTY_SECRETARY'), branch_weight * Decimal('.2')),
                 (means.get(FillerRole.BRANCH_COMMITTEE), branch_weight * Decimal('.5')),
                 (means.get(FillerRole.BRANCH_STAFF), Decimal('.2') if section == 'team' else Decimal('.4'))]
        if section == 'team':
            parts.extend([(means.get('POLITICAL_DIRECTOR'), Decimal('.16')),
                          (means.get('POLITICAL_EXECUTIVE_DEPUTY'), Decimal('.12')),
                          (means.get('POLITICAL_DEPUTY'), Decimal('.12'))])
    else:
        return None, '对象类别尚未确定，暂不计算最终分。'
    if any(value is None for value, weight in parts):
        return None, '缺少原统计表要求的评价来源或有效维度，暂不计算最终分。'
    return sum(value * weight for value, weight in parts), ''


def build_statistics(campaign):
    from .services import campaign_summary, accumulate_votes, DIMENSION_LABELS, _grade_bucket
    tasks = list(campaign.tasks.all())
    submitted = [task for task in tasks if task.status == TaskStatus.SUBMITTED]
    groups, votes, recognition = {}, {}, {}
    pending = Counter()
    dimensions = [(item['key'], item['label']) for item in TEAM_DIMENSIONS]
    # Pending targets come only from the immutable dispatch context, never draft payloads.
    for task in tasks:
        if task.form_type not in MATRIX_FORMS or task.status == TaskStatus.SUBMITTED:
            continue
        context = task.context_json or {}
        snapshots = context.get('targets') or []
        if task.form_type == FormType.ATTACHMENT_6_2 and not snapshots:
            snapshots = [{**(context.get('branch') or {}),
                          'id': (context.get('branch') or {}).get('id') or task.branch_id,
                          'name': task.branch_name_snapshot}]
        seen = set()
        for index, snapshot in enumerate(snapshots):
            section = _section(task, snapshot)
            if section not in dict(SECTIONS):
                section = 'unclassified'
            identity = str(snapshot['id']) if snapshot.get('id') is not None else f'legacy:{task.pk}:{index}'
            key = section, identity
            if key in seen:
                continue
            seen.add(key)
            pending[key] += 1
            groups.setdefault(key, {'target_id': identity, 'target_name': snapshot.get('name') or '',
                'department': snapshot.get('department') or '',
                'values': defaultdict(lambda: defaultdict(list)), 'missing': Counter(), 'expected': Counter()})
            if task.form_type in {FormType.ATTACHMENT_6_2, FormType.ATTACHMENT_7_4}:
                recognition.setdefault(key, {'target_id': f'{section}:{identity}',
                    'target_name': snapshot.get('name') or '', 'department': snapshot.get('department') or '',
                    'counts': Counter(), 'missing_count': 0})
    for task in submitted:
        if task.form_type not in MATRIX_FORMS:
            # Preserve basic counting for self evaluation, branch evaluation and interviews.
            # Legacy interview rows lack a reliable id in accumulate_votes: keep task-local.
            for index, (form, name, dimension, grade) in enumerate(accumulate_votes(task)):
                bucket = _grade_bucket(grade)
                if not bucket:
                    continue
                identity = (str(task.roster_id or task.assignee_id) if form == FormType.ATTACHMENT_2
                            else str(task.branch_id) if form == FormType.ATTACHMENT_4 and task.branch_id
                            else f'legacy:{task.pk}:{index}')
                key = form, identity, task.filler_role, dimension
                row = votes.setdefault(key, {'form_type': form, 'form_label': task.get_form_type_display(),
                    'target_id': identity, 'target_name': name, 'filler_role': task.filler_role,
                    'filler_role_label': task.get_filler_role_display(), 'dimension': DIMENSION_LABELS.get(dimension, dimension),
                    'excellent': 0, 'good': 0, 'average': 0, 'poor': 0, 'total': 0})
                row[bucket] += 1
                row['total'] += 1
            continue
        for target, snapshot, index in _targets(task):
            section = _section(task, snapshot)
            if section not in dict(SECTIONS):
                section = 'unclassified'
            identity = str(target.get('id')) if target.get('id') is not None else f'legacy:{task.pk}:{index}'
            key = section, identity
            entry = groups.setdefault(key, {'target_id': identity, 'target_name': snapshot.get('name') or target.get('name') or '',
                'department': snapshot.get('department') or target.get('department') or '',
                'values': defaultdict(lambda: defaultdict(list)), 'missing': Counter(), 'expected': Counter()})
            role = FillerRole.BRANCH_STAFF if task.form_type == FormType.ATTACHMENT_6_2 else task.filler_role
            if task.form_type == FormType.ATTACHMENT_7_4 and section in {'chief', 'deputy'} and role in {
                FillerRole.BRANCH_SECRETARY, 'BRANCH_DEPUTY_SECRETARY', FillerRole.BRANCH_COMMITTEE
            }:
                role = FillerRole.BRANCH_COMMITTEE
            for dimension, label in dimensions:
                grade = (target.get('scores') or {}).get(dimension)
                entry['expected'][dimension] += 1
                if grade not in GRADES:
                    entry['missing'][dimension] += 1
                    continue
                entry['values'][dimension][role].append(GRADES[grade])
                vote_key = task.form_type, identity, role, dimension
                row = votes.setdefault(vote_key, {'form_type': task.form_type, 'form_label': task.get_form_type_display(),
                    'target_id': identity, 'target_name': entry['target_name'], 'filler_role': role,
                    'filler_role_label': dict(FillerRole.choices).get(role, role), 'dimension': label,
                    'excellent': 0, 'good': 0, 'average': 0, 'poor': 0, 'total': 0})
                row[dict(zip(GRADES, ['excellent', 'good', 'average', 'poor']))[grade]] += 1
                row['total'] += 1
            if task.form_type in {FormType.ATTACHMENT_6_2, FormType.ATTACHMENT_7_4}:
                rec = recognition.setdefault(key, {'target_id': f'{section}:{identity}', 'target_name': entry['target_name'],
                    'department': entry['department'], 'counts': Counter(), 'missing_count': 0})
                value = target.get('recognition')
                if value in RECOGNITION:
                    rec['counts'][value] += 1
                else:
                    rec['missing_count'] += 1
    sections = {key: {'key': key, 'label': label, 'rows': []} for key, label in SECTIONS}
    for (section, identity), entry in groups.items():
        entry['pending_count'] = pending[(section, identity)]
        per_dimension = {}
        for dimension, label in dimensions:
            values = entry['values'][dimension]
            means = {role: _mean(scores) for role, scores in values.items()}
            if section == 'branch':
                means['leaders'] = _mean(values.get(FillerRole.PRISON_LEADER, []) + values.get(FillerRole.PRINCIPAL_LEADER, []))
            per_dimension[dimension] = means
            final, reason = _final(section, means)
            counts = {role: len(scores) for role, scores in values.items()}
            if section == 'branch':
                counts['leaders'] = counts.get(FillerRole.PRISON_LEADER, 0) + counts.get(FillerRole.PRINCIPAL_LEADER, 0)
            sections[section]['rows'].append(_row(entry, dimension, label, means, final, reason,
                sum(len(scores) for scores in values.values()), entry['missing'][dimension], counts))
        roles = set().union(*(means.keys() for means in per_dimension.values()))
        overall = {role: _mean([means[role] for means in per_dimension.values()])
                   for role in roles if all(means.get(role) is not None for means in per_dimension.values())}
        final, reason = _final(section, overall)
        counts = {role: sum(len(entry['values'][dim].get(role, [])) for dim, label in dimensions) for role in roles}
        if section == 'branch':
            counts['leaders'] = counts.get(FillerRole.PRISON_LEADER, 0) + counts.get(FillerRole.PRINCIPAL_LEADER, 0)
        sections[section]['rows'].append(_row(entry, 'overall', '综合分（七维平均）', overall, final, reason,
            sum(entry['expected'].values()) - sum(entry['missing'].values()), sum(entry['missing'].values()), counts))
    for key, rec in recognition.items():
        counts = rec.pop('counts')
        total = sum(counts.values())
        score = sum(Decimal(RECOGNITION[k]) * v for k, v in counts.items()) / total if total else None
        remaining = pending[key]
        sections['recognition']['rows'].append({**rec, 'dimension': '认可度', 'dimension_key': 'recognition',
            'effective_count': total, 'pending_count': remaining,
            'final_score': None if remaining or rec['missing_count'] else rounded(score), 'counts': dict(counts),
            'group_scores_text': '；'.join(f'{key} {counts[key]}票' for key in RECOGNITION),
            'reason': (f'尚有{remaining}份未提交，暂不计算最终分。' if remaining else
                       f'已提交记录缺少{rec["missing_count"]}项有效认可度评价，暂不计算最终分。' if rec['missing_count'] else
                       '按实际有效人数平均；不了解按原表计75分。')})
    return {**campaign_summary(campaign, include_progress=True), 'catalog': form_catalog(),
            'submitted_count': len(submitted), 'total_count': len(tasks), 'votes': list(votes.values()),
            'report_sections': list(sections.values()), 'note': NOTE}


def _row(entry, dimension, label, means, final, reason, count, missing, counts):
    labels = dict(FillerRole.choices)
    labels['leaders'] = '监狱领导合计'
    scores = {role: rounded(value) for role, value in means.items()}
    pending = entry['pending_count']
    if pending:
        final, reason = None, f'尚有{pending}份未提交（含草稿、退回），暂不计算最终分。'
    elif missing:
        final, reason = None, f'已提交记录缺少{missing}项有效评价，暂不计算最终分。'
    return {key: entry[key] for key in ('target_id', 'target_name', 'department')} | {
        'dimension_key': dimension, 'dimension': label, 'group_scores': scores, 'group_counts': counts,
        'group_scores_text': '；'.join(f'{labels.get(role, role)} {value if value is not None else "待收齐"}（有效{counts.get(role, 0)}项）' for role, value in scores.items()),
        'effective_count': count, 'missing_count': missing, 'pending_count': pending,
        'final_score': rounded(final), 'reason': reason}


def statistics_workbook(campaign):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    report = build_statistics(campaign)
    workbook = Workbook()
    workbook.remove(workbook.active)
    names = {'branch': '班子统计', 'chief': '正职统计', 'deputy': '副职统计', 'team': '团队负责人统计',
             'police': '警员职级统计', 'unclassified': '待确认对象类别', 'recognition': '认可度'}
    def safe(value):
        return "'" + value if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')) else value
    def append(sheet, values):
        sheet.append([safe(value) for value in values])
    for section in report['report_sections']:
        sheet = workbook.create_sheet(names[section['key']])
        sheet.append(['对象编号', '对象', '部门', '维度', '分组均分 / 认可度票数', '有效评价数', '缺失评价数', '最终分', '说明', '待提交份数'])
        for row in section['rows']:
            append(sheet, [row.get(key) for key in ['target_id', 'target_name', 'department', 'dimension',
                   'group_scores_text', 'effective_count', 'missing_count', 'final_score', 'reason', 'pending_count']])
    sheet = workbook.create_sheet('计票明细')
    sheet.append(['表单', '对象编号', '对象', '评价来源', '维度', '优', '良', '中', '差', '有效票数'])
    for row in report['votes']:
        append(sheet, [row.get(key) for key in ['form_label', 'target_id', 'target_name', 'filler_role_label',
                      'dimension', 'excellent', 'good', 'average', 'poor', 'total']])
    sheet = workbook.create_sheet('统计口径')
    append(sheet, ['批次', campaign.name])
    append(sheet, ['范围', NOTE])
    append(sheet, ['班子', '监狱领导40% + 书记20% + 民警职工40%'])
    append(sheet, ['正副职', '书记20% + (主要评价监狱领导30% + 其他监狱领导70%)*40% + 支委20% + 民警职工20%'])
    append(sheet, ['团队负责人', '(政治主任40%+常务副主任30%+副主任均分30%)*40%+(书记30%+副书记20%+支委50%)*40%+民警20%'])
    append(sheet, ['警员职级', '(书记30%+副书记20%+支委50%)*60%+民警40%'])
    append(sheet, ['缺失', '缺失数为已提交记录内缺失/无效评分，不含未提交任务；综合行有效数为七维有效评价项总数。'])
    append(sheet, ['待提交', '根据下发时的对象名单统计待填、草稿、退回任务份数；对象仍有任何待提交任务时不产生最终分。'])
    append(sheet, ['旧记录', '缺少对象编号时按任务隔离，不按姓名合并；需重新下发补齐对象编号后才能跨任务汇总。'])
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(65, max(14, max(len(str(c.value or '')) for c in column) + 2))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
