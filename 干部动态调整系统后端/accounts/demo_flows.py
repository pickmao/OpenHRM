"""Isolated, repeatable demo forms and historical department transfers."""
from datetime import timedelta
from pathlib import Path
import secrets

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from audit.models import AuditLog, AuditAction
from forms.models import (
    FormTemplate, DispatchBatch, DispatchRule, FormTask, FormSubmission,
    FormTaskAudit, OnlyOfficeDocument,
)
from orgs.models import Membership


def _workbook(relative_path, title, rows):
    """Create only in the dedicated demo directory; never replace edited files."""
    path = Path(settings.MEDIA_ROOT) / relative_path
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        book = Workbook()
        sheet = book.active
        sheet.title = '演示填报'
        sheet.append([title, '仅供系统测试，非真实人员资料'])
        sheet.append(['项目', '填报内容'])
        for row in rows:
            sheet.append(row)
        sheet.column_dimensions['A'].width = 24
        sheet.column_dimensions['B'].width = 70
        for cell in sheet[2]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1677FF')
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.freeze_panes = 'B3'
        with path.open('xb') as stream:
            book.save(stream)
    return path


def seed(ctx):
    counts = {'form_templates': 0, 'form_batches': 0, 'form_tasks': 0,
              'form_submissions': 0, 'transfer_history': 0}
    people = list(ctx.users) + [ctx.admin]
    states = ['PENDING', 'DRAFT', 'SUBMITTED', 'RETURNED', 'OVERDUE', 'SUBMITTED']
    for batch_index, label in enumerate(['月度工作填报', '季度履职总结']):
        title = ctx.prefix + label
        relative = f'forms/demo-v1/templates/template-{batch_index}.xlsx'
        _workbook(relative, title, [('姓名', ''), ('部门', ''), ('工作内容', ''), ('存在问题', ''), ('改进计划', '')])
        template = ctx.put(FormTemplate, f'flows/template/{batch_index}',
                           code=f'DEMO_V1_FORM_{batch_index + 1}', name=title,
                           description='独立演示模板，可下载或在线填写；不含真实资料。',
                           source_file=relative, source_file_name=title + '.xlsx', created_by=ctx.admin)
        closed = batch_index == 1
        batch = ctx.put(DispatchBatch, f'flows/batch/{batch_index}', name=title + '批次',
                        cycle_id='DEMO-V1', stage_code=f'DEMO-{batch_index + 1}',
                        status='CLOSED' if closed else 'PUBLISHED',
                        deadline_at=ctx.now + timedelta(days=-3 if closed else 14),
                        published_at=ctx.now - timedelta(days=15),
                        closed_at=ctx.now - timedelta(days=1) if closed else None,
                        created_by=ctx.admin)
        assignees = [ctx.users[0], ctx.users[4], ctx.admin] if closed else people
        rule = ctx.put(DispatchRule, f'flows/rule/{batch_index}', batch=batch, template=template,
                       receiver_type='USER', receiver_expr_json={'user_ids': [str(u.id) for u in assignees]},
                       target_type='USER', created_by=ctx.admin)
        for index, user in enumerate(assignees):
            state = 'SUBMITTED' if closed else ('PENDING' if user == ctx.admin else states[index % len(states)])
            unit = ctx.departments[ctx.users.index(user) // 2] if user != ctx.admin else ctx.departments[0]
            final = state == 'SUBMITTED'
            submitted_at = ctx.now - timedelta(days=4 if closed else 1) if final else None
            task_key = f'flows/task/{batch_index}/{index}'
            task = ctx.put(FormTask, task_key, batch=batch, template=template, template_version='1.0',
                           assignee_type='USER', assignee_id=user.id,
                           assignee_name_snapshot=user.real_name or user.username,
                           org_unit_snapshot=unit, target_type='USER', target_id=user.id,
                           status=state, deadline_at=ctx.now - timedelta(days=2) if state == 'OVERDUE' else batch.deadline_at,
                           submitted_at=submitted_at,
                           context_json={'demo': True, 'dispatch_rule_id': str(rule.id), 'source_file_name': title + '.xlsx'})
            ctx.put(FormTaskAudit, task_key + '/dispatch', task=task, action='DISPATCH', operator=ctx.admin,
                    after_json={'batch_id': str(batch.id)}, reason='演示数据初始化')
            filled = state in {'DRAFT', 'SUBMITTED', 'RETURNED'}
            payload = {'姓名': user.real_name, '部门': unit.name,
                       '工作内容': '【演示】完成本月业务培训和台账整理。',
                       '存在问题': '【演示】跨部门协作时效有待提升。',
                       '改进计划': '【演示】每周跟踪事项并反馈进展。'} if filled else {}
            if filled:
                ctx.put(FormSubmission, task_key + '/submission', task=task, version=1, is_final=final,
                        payload_json=payload, submitted_by=user, submitted_at=submitted_at,
                        return_reason='【演示】请补充改进计划的完成时间。' if state == 'RETURNED' else '',
                        returned_by=ctx.admin if state == 'RETURNED' else None,
                        returned_at=ctx.now if state == 'RETURNED' else None)
                ctx.put(FormTaskAudit, task_key + '/save', task=task,
                        action='SUBMIT' if final else ('RETURN' if state == 'RETURNED' else 'SAVE_DRAFT'),
                        operator=ctx.admin if state == 'RETURNED' else user,
                        after_json={'status': state}, reason='演示填报状态')
                counts['form_submissions'] += 1
            task_file = f'forms/demo-v1/tasks/{task.id}.xlsx'
            path = _workbook(task_file, title, list(payload.items()) if filled else [
                ('姓名', user.real_name), ('部门', unit.name), ('工作内容', ''), ('存在问题', ''), ('改进计划', '')])
            ctx.put(OnlyOfficeDocument, task_key + '/document', task=task,
                    original_name=title + '.xlsx', file_ext='xlsx',
                    mime_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    file_size=path.stat().st_size, storage_path=str(path), access_token=secrets.token_hex(32),
                    file=task_file, file_name=title + '.xlsx', document_key=f'form-{task.id}-v1',
                    last_saved_at=ctx.now if filled else None, created_by=user)
            counts['form_tasks'] += 1
        batch.refresh_stats()
        counts['form_templates'] += 1
        counts['form_batches'] += 1

    # Historical moves end in the current department, so the shared demo roster
    # and active memberships remain consistent without mutating current records.
    for index in [1, 5, 9]:
        user = ctx.users[index]
        target = ctx.departments[index // 2]
        source = ctx.departments[(index // 2 + 1) % len(ctx.departments)]
        current = Membership.objects.filter(user=user, unit=target, effective_to__isnull=True).first()
        effective_date = (current.effective_from if current else None) or (ctx.now - timedelta(days=30)).date()
        previous = ctx.put(Membership, f'flows/old-membership/{index}', user=user, unit=source,
                           is_primary=True, position='【演示】历史岗位', effective_from=effective_date - timedelta(days=180),
                           effective_to=effective_date - timedelta(days=1), created_by=ctx.admin)
        if current is None:
            current = ctx.put(Membership, f'flows/current-membership/{index}', user=user, unit=target,
                              is_primary=True, effective_from=effective_date, created_by=ctx.admin)
        ctx.put(AuditLog, f'flows/transfer/{index}', actor=ctx.admin, action=AuditAction.UPDATE_ORG,
                target_type='Membership', target_id=current.id,
                context={'from_dept': str(source.id), 'to_dept': str(target.id),
                         'reason': '【演示】岗位轮训后调整至现部门。', 'demo': True,
                         'effective_date': effective_date.isoformat(), 'user_id': str(user.id),
                         'previous_membership_id': str(previous.id)})
        counts['transfer_history'] += 1
    return counts
