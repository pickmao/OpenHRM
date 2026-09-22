from collections import defaultdict
from io import BytesIO

from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import PermissionDenied, ValidationError

from accounts.models import User
from cadres.models import PersonnelRoster
from orgs.models import Membership, OrgUnit, UnitType

from .models import (
    POST_CATEGORIES, POST_CATEGORY_MAP, RECOMMENDER_CATEGORIES, CampaignStatus,
    RecommendationCampaign, RecommendationNomination, RecommendationTask, TaskStatus,
    default_slot_config, normalize_slot_config,
)


RULES_TEXT = (
    '岗位类别与空位数以《党支部优秀干部推荐表》为准：正科级领导干部、副科级领导干部、'
    '工作团队负责人（正职）、工作团队负责人（副职）、警员职级警察；模板默认每类 1 人，'
    '管理员下发时可按每个岗位类别分别设置推荐人数。'
    '候选人从花名册按填报人所在支部 / 部门 / 监区过滤，禁止手填姓名。'
    '同一人默认不能在同一张表的多个岗位重复推荐。'
    '本人可否推荐本人由管理员配置，默认不允许。'
    '每空可填写 1 人，无则留空。'
)


def _active_memberships(user=None):
    today = timezone.localdate()
    queryset = Membership.objects.filter(unit__is_active=True).filter(
        Q(effective_from__isnull=True) | Q(effective_from__lte=today),
        Q(effective_to__isnull=True) | Q(effective_to__gte=today),
    )
    if user is not None:
        queryset = queryset.filter(user=user)
    return queryset.select_related('unit', 'unit__parent')


def primary_unit(user):
    membership = _active_memberships(user).order_by('-is_primary', 'unit__sort_order', 'id').first()
    return membership.unit if membership else None


def find_branch(unit):
    current = unit
    while current:
        if current.unit_type == UnitType.BRANCH:
            return current
        current = current.parent
    return None


def collect_scope(user):
    """按填报人组织归属确定花名册过滤范围。"""
    unit = primary_unit(user)
    if not unit:
        return {
            'unit': None,
            'branch': None,
            'branch_name': '',
            'org_unit_name': '',
            'department_names': [],
            'scope_label': '未关联组织',
            'fallback': 'none',
        }

    branch = find_branch(unit)
    if unit.unit_type == UnitType.BRANCH:
        branch = unit
    names = []
    if branch:
        names.append(branch.name)
        for child in branch.get_descendants():
            names.append(child.name)
        fallback = 'branch'
        scope_label = branch.name
    else:
        names.append(unit.name)
        for child in unit.get_descendants():
            names.append(child.name)
        fallback = 'department'
        scope_label = unit.name
        branch = None

    unique_names = []
    seen = set()
    for name in names:
        text = (name or '').strip()
        if text and text not in seen:
            seen.add(text)
            unique_names.append(text)
    return {
        'unit': unit,
        'branch': branch,
        'branch_name': branch.name if branch else '',
        'org_unit_name': unit.name,
        'department_names': unique_names,
        'scope_label': scope_label,
        'fallback': fallback,
    }


def _descendant_unit_ids(units):
    ids = set()
    for unit in units:
        ids.add(unit.id)
        for child in unit.get_descendants():
            ids.add(child.id)
    return ids


def resolve_recipients(receiver_type, expr):
    expr = expr or {}
    if receiver_type == 'USER':
        if expr.get('all_users'):
            users = User.objects.filter(is_active=True)
        else:
            ids = expr.get('user_ids') or expr.get('ids') or []
            users = User.objects.filter(id__in=ids, is_active=True)
        return list(users.order_by('real_name', 'username'))

    if receiver_type == 'ORG_UNIT':
        ids = expr.get('org_unit_ids') or expr.get('ids') or []
        units = list(OrgUnit.objects.filter(id__in=ids, is_active=True))
        unit_ids = _descendant_unit_ids(units)
        if not unit_ids:
            return []
        return list(
            User.objects.filter(
                is_active=True,
                memberships__unit_id__in=unit_ids,
                memberships__effective_to__isnull=True,
            ).distinct().order_by('real_name', 'username')
        )

    if receiver_type == 'ORG_ROLE':
        role_codes = expr.get('role_codes') or ([expr['role_code']] if expr.get('role_code') else [])
        unit_ids = expr.get('org_unit_ids') or []
        users = User.objects.filter(
            is_active=True,
            user_roles__role__is_active=True,
            user_roles__role__code__in=role_codes,
        )
        if unit_ids:
            units = list(OrgUnit.objects.filter(id__in=unit_ids, is_active=True))
            users = users.filter(
                memberships__unit_id__in=_descendant_unit_ids(units),
                memberships__effective_to__isnull=True,
            )
        return list(users.distinct().order_by('real_name', 'username'))

    raise ValidationError({'receiver_type': '不支持的接收人类型'})


def position_blob(person):
    return ''.join([
        person.position or '',
        person.position_category or '',
        person.position_rank or '',
        person.position_level or '',
        person.police_rank or '',
        person.police_title or '',
    ])


def roster_matches_post(person, key):
    blob = position_blob(person)
    position = person.position or ''
    category = person.position_category or ''
    if key == 'section_chief':
        return '正科' in blob
    if key == 'deputy_section_chief':
        return '副科' in blob
    if key == 'team_lead':
        return category == '监区工作团队正职' or ('团队' in blob and '副' not in position and '副职' not in blob)
    if key == 'deputy_team_lead':
        return category == '监区工作团队副职' or ('团队' in blob and ('副' in position or '副职' in blob))
    if key == 'police_officer':
        return '警员' in blob or category not in {
            '领导职务', '内定领导职务', '监区工作团队正职', '监区工作团队副职',
        }
    return True


def format_position(person):
    parts = [person.position, person.position_rank or person.position_level or person.police_rank]
    return ' / '.join(part for part in parts if part)


def _is_self(person, user):
    real_name = (user.real_name or '').strip()
    return bool(real_name) and person.name.strip() == real_name


def roster_queryset_for_scope(department_names):
    if not department_names:
        return PersonnelRoster.objects.none()
    query = Q()
    for name in department_names:
        query |= Q(department=name)
    return PersonnelRoster.objects.filter(query).order_by('department', 'serial_number', 'name')


def serialize_candidate(person):
    return {
        'id': str(person.id),
        'name': person.name,
        'department': person.department,
        'position': person.position,
        'position_rank': person.position_rank or person.position_level or person.police_rank,
        'position_label': format_position(person),
        'label': f'{person.name}（{person.department} · {format_position(person) or "职务未登记"}）',
    }


def candidates_for_task(task):
    campaign = task.campaign
    user = task.assignee
    scope = collect_scope(user)
    people = list(roster_queryset_for_scope(scope['department_names']))
    if not campaign.allow_self_recommend:
        people = [person for person in people if not _is_self(person, user)]

    categories = []
    for item in POST_CATEGORIES:
        slots = int((campaign.slot_config or {}).get(item['key'], item['default_slots']))
        if slots <= 0:
            continue
        matched = [person for person in people if roster_matches_post(person, item['key'])]
        used_fallback = not matched and bool(people)
        source = people if used_fallback else matched
        categories.append({
            **item,
            'slots': slots,
            'position_filtered': not used_fallback,
            'candidates': [serialize_candidate(person) for person in source],
        })
    return {
        'scope_label': scope['scope_label'],
        'branch_name': scope['branch_name'] or task.branch_name_snapshot,
        'org_unit_name': scope['org_unit_name'] or task.org_unit_name_snapshot,
        'department_names': scope['department_names'],
        'fallback': scope['fallback'],
        'allow_self_recommend': campaign.allow_self_recommend,
        'post_categories': categories,
        'candidate_count': len(people),
    }


def campaign_summary(campaign, *, include_progress=False):
    payload = {
        'id': str(campaign.id),
        'name': campaign.name,
        'description': campaign.description,
        'status': campaign.status,
        'status_display': campaign.get_status_display(),
        'slot_config': campaign.slot_config or default_slot_config(),
        'allow_self_recommend': campaign.allow_self_recommend,
        'receiver_type': campaign.receiver_type,
        'deadline_at': campaign.deadline_at,
        'published_at': campaign.published_at,
        'closed_at': campaign.closed_at,
        'is_open': campaign.is_open(),
        'post_categories': POST_CATEGORIES,
        'recommender_categories': RECOMMENDER_CATEGORIES,
        'rules_text': RULES_TEXT,
    }
    if include_progress:
        total = campaign.tasks.count()
        submitted = campaign.tasks.filter(status=TaskStatus.SUBMITTED).count()
        payload['progress'] = {
            'total': total,
            'submitted': submitted,
            'completion_rate': round(submitted * 100 / total, 1) if total else 0,
        }
    return payload


def preview_dispatch(payload):
    users = resolve_recipients(payload.get('receiver_type'), payload.get('receiver_expr_json') or {})
    tasks = []
    for user in users:
        scope = collect_scope(user)
        tasks.append({
            'assignee_id': str(user.id),
            'assignee_name': user.real_name or user.username,
            'org_unit': scope['org_unit_name'],
            'branch_name': scope['branch_name'],
            'scope_label': scope['scope_label'],
        })
    return {'summary': {'total': len(tasks)}, 'tasks': tasks}


@transaction.atomic
def create_campaign(data, creator):
    users = resolve_recipients(data['receiver_type'], data.get('receiver_expr_json') or {})
    if not users:
        raise ValidationError({'receiver_expr_json': '按当前范围未解析到任何填报人。'})
    now = timezone.now()
    campaign = RecommendationCampaign.objects.create(
        name=data['name'].strip(),
        description=(data.get('description') or '').strip(),
        status=CampaignStatus.PUBLISHED,
        slot_config=normalize_slot_config(data.get('slot_config') or default_slot_config()),
        allow_self_recommend=bool(data.get('allow_self_recommend', False)),
        receiver_type=data['receiver_type'],
        receiver_expr_json=data.get('receiver_expr_json') or {},
        deadline_at=data['deadline_at'],
        published_at=now,
        created_by=creator,
    )
    RecommendationTask.objects.bulk_create([
        RecommendationTask(
            campaign=campaign,
            assignee=user,
            assignee_name_snapshot=user.real_name or user.username,
            org_unit_name_snapshot=scope['org_unit_name'],
            branch_name_snapshot=scope['branch_name'] or scope['org_unit_name'],
        )
        for user in users
        for scope in [collect_scope(user)]
    ])
    return campaign


@transaction.atomic
def close_campaign(campaign):
    campaign = RecommendationCampaign.objects.select_for_update().get(id=campaign.id)
    if campaign.status != CampaignStatus.PUBLISHED:
        raise ValidationError('只有填写中的活动可以关闭。')
    campaign.status = CampaignStatus.CLOSED
    campaign.closed_at = timezone.now()
    campaign.save(update_fields=['status', 'closed_at', 'updated_at'])
    return campaign


def _serialize_nominations(task):
    return [
        {
            'id': str(item.id),
            'post_category': item.post_category,
            'slot_index': item.slot_index,
            'roster_id': str(item.roster_id),
            'name': item.name_snapshot,
            'department': item.roster.department if item.roster_id else item.department_snapshot,
            'department_snapshot': item.department_snapshot,
            'position_label': item.position_snapshot,
        }
        for item in task.nominations.select_related('roster').order_by('post_category', 'slot_index')
    ]


def serialize_task(task, *, include_candidates=False):
    campaign = task.campaign
    payload = {
        'id': str(task.id),
        'campaign_id': str(campaign.id),
        'campaign_name': campaign.name,
        'status': task.status,
        'status_display': task.get_status_display(),
        'recommender_category': task.recommender_category,
        'assignee_name': task.assignee_name_snapshot,
        'org_unit_name': task.org_unit_name_snapshot,
        'branch_name': task.branch_name_snapshot,
        'deadline_at': campaign.deadline_at,
        'is_open': campaign.is_open() and task.status != TaskStatus.SUBMITTED,
        'submitted': task.status == TaskStatus.SUBMITTED,
        'allow_self_recommend': campaign.allow_self_recommend,
        'slot_config': campaign.slot_config or default_slot_config(),
        'post_categories': POST_CATEGORIES,
        'recommender_categories': RECOMMENDER_CATEGORIES,
        'rules_text': RULES_TEXT,
        'nominations': _serialize_nominations(task),
    }
    if include_candidates:
        payload['candidates'] = candidates_for_task(task)
    return payload


def _validate_nominations(task, nominations, *, require_category=False, recommender_category=''):
    campaign = task.campaign
    slot_config = campaign.slot_config or default_slot_config()
    if require_category and not recommender_category:
        raise ValidationError({'recommender_category': '请选择人员类别：党支部书记、党支部支委或党支部全体民警职工。'})

    seen_slots = set()
    seen_people = set()
    roster_ids = []
    cleaned = []
    for item in nominations or []:
        key = item['post_category']
        slot_index = item['slot_index']
        roster_id = item['roster_id']
        limit = int(slot_config.get(key, 0))
        if limit <= 0:
            raise ValidationError({'nominations': f'{POST_CATEGORY_MAP[key]["label"]} 本轮未开放推荐。'})
        if slot_index >= limit:
            raise ValidationError({'nominations': f'{POST_CATEGORY_MAP[key]["label"]} 最多推荐 {limit} 人。'})
        slot_key = (key, slot_index)
        if slot_key in seen_slots:
            raise ValidationError({'nominations': '同一岗位空位不能重复选择。'})
        if roster_id in seen_people:
            raise ValidationError({'nominations': '同一人不能在同一张表的多个岗位重复推荐。'})
        seen_slots.add(slot_key)
        seen_people.add(roster_id)
        roster_ids.append(roster_id)
        cleaned.append(item)

    people = {person.id: person for person in PersonnelRoster.objects.filter(id__in=roster_ids)}
    missing = [str(item_id) for item_id in roster_ids if item_id not in people]
    if missing:
        raise ValidationError({'nominations': '存在无效的花名册人员，请重新选择。'})

    allowed = {
        item['id']
        for category in candidates_for_task(task)['post_categories']
        for item in category['candidates']
    }
    for item in cleaned:
        person = people[item['roster_id']]
        if not campaign.allow_self_recommend and _is_self(person, task.assignee):
            raise ValidationError({'nominations': '本活动不允许推荐本人。'})
        if str(person.id) not in allowed:
            raise ValidationError({'nominations': f'{person.name} 不在当前支部/部门/监区可选范围内。'})
        item['person'] = person
    return cleaned


def _replace_nominations(task, cleaned, recommender_category, *, final=False):
    task.nominations.all().delete()
    RecommendationNomination.objects.bulk_create([
        RecommendationNomination(
            task=task,
            post_category=item['post_category'],
            slot_index=item['slot_index'],
            roster=item['person'],
            name_snapshot=item['person'].name,
            department_snapshot=item['person'].department,
            position_snapshot=format_position(item['person']),
        )
        for item in cleaned
    ])
    task.recommender_category = recommender_category or ''
    if final:
        task.status = TaskStatus.SUBMITTED
        task.submitted_at = timezone.now()
    else:
        task.status = TaskStatus.DRAFT
        task.submitted_at = None
    task.save(update_fields=['recommender_category', 'status', 'submitted_at', 'updated_at'])
    return task


def _ensure_editable(task, user):
    if task.assignee_id != user.id:
        raise PermissionDenied('无权处理该推荐任务')
    if task.status == TaskStatus.SUBMITTED:
        raise ValidationError('该推荐表已提交，不能再修改。')
    if not task.campaign.is_open():
        raise ValidationError('活动未在填报期内，不能提交。')


@transaction.atomic
def save_draft(task, user, data):
    task = RecommendationTask.objects.select_for_update().select_related('campaign', 'assignee').get(id=task.id)
    _ensure_editable(task, user)
    cleaned = _validate_nominations(
        task, data.get('nominations') or [],
        recommender_category=data.get('recommender_category') or '',
    )
    return _replace_nominations(task, cleaned, data.get('recommender_category') or '', final=False)


@transaction.atomic
def submit_task(task, user, data):
    task = RecommendationTask.objects.select_for_update().select_related('campaign', 'assignee').get(id=task.id)
    _ensure_editable(task, user)
    cleaned = _validate_nominations(
        task, data.get('nominations') or [],
        require_category=True,
        recommender_category=data.get('recommender_category') or '',
    )
    return _replace_nominations(task, cleaned, data.get('recommender_category') or '', final=True)


def build_statistics(campaign):
    submitted = list(
        campaign.tasks.filter(status=TaskStatus.SUBMITTED).prefetch_related('nominations__roster')
    )
    grouped = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda: {
        'name': '', 'department': '', 'position_label': '', 'roster_id': '', 'vote_count': 0, 'person_ids': set(),
    })))
    submitted_by_branch = defaultdict(int)
    for task in submitted:
        branch_name = task.branch_name_snapshot or task.org_unit_name_snapshot or '未划分支部'
        submitted_by_branch[branch_name] += 1
        for item in task.nominations.all():
            bucket = grouped[branch_name][item.post_category][item.roster_id]
            bucket['name'] = item.roster.name if item.roster_id else item.name_snapshot
            bucket['department'] = item.roster.department if item.roster_id else item.department_snapshot
            bucket['position_label'] = item.position_snapshot
            bucket['roster_id'] = str(item.roster_id)
            bucket['vote_count'] += 1
            bucket['person_ids'].add(task.id)

    branches = []
    all_branch_names = sorted(set(submitted_by_branch) | set(grouped))
    if not all_branch_names:
        all_branch_names = []
    for branch_name in all_branch_names:
        rows = []
        for category in POST_CATEGORIES:
            people = grouped[branch_name].get(category['key'], {})
            ordered = sorted(people.values(), key=lambda item: (-item['vote_count'], item['name']))
            if not ordered:
                rows.append({
                    'post_category': category['key'],
                    'post_category_label': category['stats_label'],
                    'name': '',
                    'department': '',
                    'position_label': '',
                    'roster_id': '',
                    'person_count': None,
                    'vote_count': None,
                })
                continue
            for person in ordered:
                rows.append({
                    'post_category': category['key'],
                    'post_category_label': category['stats_label'],
                    'name': person['name'],
                    'department': person['department'],
                    'position_label': person['position_label'],
                    'roster_id': person['roster_id'],
                    'person_count': len(person['person_ids']),
                    'vote_count': person['vote_count'],
                })
        branches.append({
            'branch_name': branch_name,
            'title': f'{branch_name}党支部优秀干部统计表' if '支部' not in branch_name else f'{branch_name}优秀干部统计表',
            'submitted_count': submitted_by_branch.get(branch_name, 0),
            'rows': rows,
        })
    return {
        'campaign': campaign_summary(campaign, include_progress=True),
        'columns': ['人员类别', '姓名', '职务职级', '统计人次', '推荐票数'],
        'submitted_count': len(submitted),
        'total_count': campaign.tasks.count(),
        'branches': branches,
    }


def export_statistics_xlsx(campaign):
    stats = build_statistics(campaign)
    workbook = Workbook()
    thin = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )
    header_font = Font(bold=True, size=14)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    if stats['branches']:
        first = True
        for branch in stats['branches']:
            sheet = workbook.active if first else workbook.create_sheet()
            first = False
            sheet.title = (branch['branch_name'] or '未划分支部')[:31]
            _write_stats_sheet(sheet, branch, thin, header_font, center)
    else:
        sheet = workbook.active
        sheet.title = '统计表'
        _write_stats_sheet(sheet, {
            'title': f'{campaign.name}统计表',
            'rows': [],
            'submitted_count': 0,
        }, thin, header_font, center)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _write_stats_sheet(sheet, branch, thin, header_font, center):
    title = branch.get('title') or f'{branch.get("branch_name", "")}优秀干部统计表'
    sheet.merge_cells('A1:E1')
    sheet['A1'] = title
    sheet['A1'].font = header_font
    sheet['A1'].alignment = center
    headers = ['人员类别', '姓名', '职务职级', '统计人次', '推荐票数']
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=2, column=index, value=header)
        cell.alignment = center
        cell.border = thin
        cell.font = Font(bold=True)

    rows = branch.get('rows') or []
    if not rows:
        rows = [{
            'post_category_label': item['stats_label'], 'name': '', 'position_label': '',
            'person_count': None, 'vote_count': None,
        } for item in POST_CATEGORIES]

    start_row = 3
    for offset, item in enumerate(rows):
        row_number = start_row + offset
        values = [
            item.get('post_category_label') or '',
            item.get('name') or '',
            item.get('position_label') or '',
            item.get('person_count') if item.get('person_count') is not None else '',
            item.get('vote_count') if item.get('vote_count') is not None else '',
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column, value=value)
            cell.alignment = center
            cell.border = thin

    merge_start = start_row
    for index, item in enumerate(rows):
        row_number = start_row + index
        next_label = rows[index + 1].get('post_category_label') if index + 1 < len(rows) else None
        if next_label != item.get('post_category_label'):
            if row_number > merge_start:
                sheet.merge_cells(start_row=merge_start, start_column=1, end_row=row_number, end_column=1)
            merge_start = row_number + 1

    footer = start_row + len(rows) + 1
    sheet.merge_cells(start_row=footer, start_column=1, end_row=footer, end_column=2)
    sheet.cell(row=footer, column=1, value='监票人：')
    sheet.merge_cells(start_row=footer, start_column=4, end_row=footer, end_column=5)
    sheet.cell(row=footer, column=4, value='监票人：')
    sheet.cell(row=footer + 1, column=1, value=f"已填报人次：{branch.get('submitted_count', 0)}")

    widths = [22, 14, 28, 12, 12]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 28
