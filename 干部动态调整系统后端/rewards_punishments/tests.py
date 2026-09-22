from io import BytesIO
from tempfile import TemporaryDirectory

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook
from rest_framework import status
from rest_framework.test import APIClient

from accounts.models import User
from .models import RewardRecipientType, RewardRecord


class RewardImportApiTests(TestCase):
    def setUp(self):
        self.media_dir = TemporaryDirectory()
        self.settings_override = override_settings(MEDIA_ROOT=self.media_dir.name)
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.addCleanup(self.media_dir.cleanup)
        self.user = User.objects.create_user(username='reward_admin', password='password', real_name='奖励管理员', is_superuser=True)
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    @staticmethod
    def _reward_workbook():
        workbook = Workbook()
        individual = workbook.active
        individual.title = '个人'
        individual.append(['序号', '姓名', '级别', '批准年度', '奖励情况', '批准时间', '文号', '备注'])
        individual.append([0, '这一行绝对不能删', '这一行绝对不能删', '这一行绝对不能删', '这一行绝对不能删', '', '', ''])
        individual.append([1, '张三', '厅局级', 2025, '年度考核优秀给予嘉奖', '2025.09.30', '粤狱发〔2025〕1号', '测试备注'])
        collective = workbook.create_sheet('集体')
        collective.append(['获奖单位', '批准年度', '授奖等级', '荣誉称号', '获奖时间', '授奖文件名号', '备注'])
        collective.append(['政治处', 2024, '监狱', '先进集体', '2024.06.25', '肇狱党〔2024〕15号', ''])
        collective.append(['', '2025年', '厅局级', '集体三等功', '2025.03.11', '粤司发〔2025〕24号', ''])
        historical_collective = workbook.create_sheet('历史集体奖励')
        historical_collective.append(['序号', '授奖等级', '获奖时间', '荣誉称号', '获奖单位', '授奖文件名号', '备注'])
        historical_collective.append([1, '司法部', '2011.08.15', '全国先进单位', '怀集监狱', '粤司办〔2011〕182号', ''])
        thirty_year_list = workbook.create_sheet('从事监狱工作三十年申报名单')
        thirty_year_list.append(['2025', '签收'])
        thirty_year_list.append(['李四', '2025.07.03签收'])
        workbook.create_sheet('说明').append(['本页不是奖励台账'])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def test_upload_stores_personal_and_collective_records_and_exposes_statistics(self):
        upload = SimpleUploadedFile(
            '监狱个人及集体奖励汇总表.xlsx',
            self._reward_workbook(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

        response = self.client.post('/api/rewards/reward-files/upload-excel/', {'file': upload}, format='multipart')

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['total_records'], 5)
        self.assertEqual(response.data['record_counts'], {'individual': 2, 'collective': 3})
        personal = RewardRecord.objects.get(recipient_name='张三')
        self.assertEqual(personal.recipient_type, RewardRecipientType.INDIVIDUAL)
        self.assertEqual(personal.approval_year, 2025)
        self.assertEqual(str(personal.approval_date), '2025-09-30')
        self.assertEqual(personal.source_sheet, '个人')
        self.assertEqual(personal.source_row, 3)

        filtered = self.client.get('/api/rewards/reward-records/', {'recipient_type': 'COLLECTIVE'})
        self.assertEqual(filtered.status_code, status.HTTP_200_OK)
        self.assertEqual(filtered.data['count'], 3)
        self.assertEqual(filtered.data['results'][0]['recipient_name'], '政治处')

        statistics = self.client.get('/api/rewards/reward-records/statistics/')
        self.assertEqual(statistics.status_code, status.HTTP_200_OK)
        self.assertEqual(statistics.data['total'], 5)
        self.assertEqual(statistics.data['individual_count'], 2)
        self.assertEqual(statistics.data['collective_count'], 3)
